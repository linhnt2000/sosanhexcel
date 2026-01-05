import polars as pl
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

def compare_excel_v3(file1, file2, key_column, output_file="compare_result_v3.xlsx"):
    # =========================
    # 1. Đọc dữ liệu với Polars (nhanh + ép TEXT)
    # =========================
    df1 = pl.read_excel(file1)
    df2 = pl.read_excel(file2)

    if key_column not in df1.columns or key_column not in df2.columns:
        raise ValueError(f"Cột khóa '{key_column}' không tồn tại.")

    # Ép về string (an toàn)
    df1 = df1.with_columns(pl.col(key_column).cast(pl.Utf8))
    df2 = df2.with_columns(pl.col(key_column).cast(pl.Utf8))

    # =========================
    # 2. Xác định trạng thái 4 trường hợp
    # =========================
    keys1 = df1[key_column].to_list()
    keys2 = df2[key_column].to_list()

    set1 = set(keys1)
    set2 = set(keys2)

    only1 = set1 - set2
    only2 = set2 - set1
    both = set1 & set2

    # Dòng duplicate >1 lần trong file2
    from collections import Counter
    count2 = Counter(keys2)
    duplicate_in_file2 = {k for k, v in count2.items() if v > 1} & set1

    # =========================
    # 3. Gán trạng thái
    # =========================
    def assign_status(df, file_index):
        status_list = []
        for val in df[key_column].to_list():
            if file_index == 1:
                if val in duplicate_in_file2:
                    status_list.append("blue")
                elif val in only1:
                    status_list.append("red")
                elif val in both:
                    status_list.append("green")
                else:
                    status_list.append("")
            else:  # file_index == 2
                if val in only2:
                    status_list.append("yellow")
                elif val in duplicate_in_file2:
                    status_list.append("blue")
                elif val in both:
                    status_list.append("green")
                else:
                    status_list.append("")
        return df.with_columns(pl.Series("__status__", status_list))

    df1 = assign_status(df1, 1)
    df2 = assign_status(df2, 2)

    # =========================
    # 4. Tạo workbook và sheet
    # =========================
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "File1"
    ws2 = wb.create_sheet("File2")

    def write_sheet(ws, df):
        # Viết header
        for c, col_name in enumerate(df.columns[:-1], start=1):  # bỏ cột __status__
            ws.cell(row=1, column=c, value=col_name)

        # Viết dữ liệu + gán màu
        color_map = {
            "green": "C6EFCE",
            "red": "FFC7CE",
            "blue": "ADD8E6",
            "yellow": "FFFACD"
        }

        for r, row in enumerate(df.to_dicts(), start=2):
            status = row["__status__"]
            for c, col_name in enumerate(df.columns[:-1], start=1):
                ws.cell(row=r, column=c, value=row[col_name])
                if status in color_map:
                    ws.cell(row=r, column=c).fill = PatternFill(
                        start_color=color_map[status],
                        end_color=color_map[status],
                        fill_type="solid"
                    )

    write_sheet(ws1, df1)
    write_sheet(ws2, df2)

    # Căn rộng cột tự động
    for ws in [ws1, ws2]:
        for col_cells in ws.columns:
            max_len = max((len(str(cell.value)) if cell.value is not None else 0) for cell in col_cells)
            ws.column_dimensions[col_cells[0].column_letter].width = max_len + 2

    wb.save(output_file)

    print(f"✅ File '{output_file}' đã được tạo thành công với màu phân loại:")
    print("🟩 green = có trong cả 2 file")
    print("🟥 red   = chỉ có trong file 1")
    print("🟦 blue  = có trong file 1 & xuất hiện >1 lần trong file 2")
    print("🟨 yellow= chỉ có trong file 2")

# =========================
# Ví dụ chạy
# =========================
compare_excel_v3("payments-Copy.xlsx", "transactions - Copy.xlsx", key_column="Payments_id")
