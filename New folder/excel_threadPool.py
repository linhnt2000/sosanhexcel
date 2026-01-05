import pandas as pd
from concurrent.futures import ThreadPoolExecutor
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment

def read_excel_parallel(file1, file2):
    """Đọc 2 file Excel song song bằng threadpool."""
    print("🔄 Đang đọc dữ liệu song song bằng ThreadPoolExecutor...")
    with ThreadPoolExecutor(max_workers=2) as executor:
        f1 = executor.submit(pd.read_excel, file1)
        f2 = executor.submit(pd.read_excel, file2)
        return f1.result(), f2.result()


def compare_excel_fast(file1, file2, column_name, output_file="compare_result.xlsx"):
    # --- Đọc dữ liệu ---
    df1, df2 = read_excel_parallel(file1, file2)

    # --- Kiểm tra cột tồn tại ---
    if column_name not in df1.columns or column_name not in df2.columns:
        raise ValueError(f"Cột '{column_name}' không tồn tại trong một trong hai file.")

    # --- Chuyển cột khóa về dạng chuỗi ---
    df1[column_name] = df1[column_name].astype(str)
    df2[column_name] = df2[column_name].astype(str)

    # --- Tập dữ liệu ---
    set1 = set(df1[column_name])
    set2 = set(df2[column_name])
    common_keys = set1 & set2
    only_in_file1 = set1 - set2
    only_in_file2 = set2 - set1
    count_in_file2 = df2[column_name].value_counts()
    duplicated_in_file2 = set(count_in_file2[count_in_file2 > 1].index) & set1

    # --- Gán màu ---
    def label_rows(df, file_index):
        df["__status__"] = ""
        for i, val in df[column_name].items():
            if file_index == 1:
                if val in only_in_file1:
                    df.loc[i, "__status__"] = "red"
                elif val in duplicated_in_file2:
                    df.loc[i, "__status__"] = "blue"
                elif val in common_keys:
                    df.loc[i, "__status__"] = "green"
            else:
                if val in only_in_file2:
                    df.loc[i, "__status__"] = "yellow"
                elif val in duplicated_in_file2:
                    df.loc[i, "__status__"] = "blue"
                elif val in common_keys:
                    df.loc[i, "__status__"] = "green"
        return df

    df1 = label_rows(df1, 1)
    df2 = label_rows(df2, 2)

    # --- Ghi ra Excel ---
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        df1.to_excel(writer, index=False, sheet_name="Sheet_File1")
        df2.to_excel(writer, index=False, sheet_name="Sheet_File2")

    # --- Tô màu ---
    wb = load_workbook(output_file)
    fills = {
        "green":  PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        "red":    PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        "blue":   PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid"),
        "yellow": PatternFill(start_color="FFFACD", end_color="FFFACD", fill_type="solid"),
    }

    for sheet_name in ["Sheet_File1", "Sheet_File2"]:
        ws = wb[sheet_name]
        status_col = ws.max_column
        header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        header_font = Font(bold=True)
        header_align = Alignment(horizontal="center", vertical="center")

        # Tô header
        for c in range(1, ws.max_column + 1):
            ws.cell(row=1, column=c).fill = header_fill
            ws.cell(row=1, column=c).font = header_font
            ws.cell(row=1, column=c).alignment = header_align

        # Tô màu dòng
        for r in range(2, ws.max_row + 1):
            color_key = ws.cell(row=r, column=status_col).value
            if color_key in fills:
                for c in range(1, ws.max_column + 1):
                    ws.cell(row=r, column=c).fill = fills[color_key]

        ws.delete_cols(status_col)

        # Căn độ rộng cột
        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value else 0) for c in col)
            ws.column_dimensions[col[0].column_letter].width = max_len + 2

    wb.save(output_file)
    print(f"✅ Đã tạo '{output_file}' (song song bằng ThreadPoolExecutor).")
    print("🟩 Xanh lá: có trong cả hai file")
    print("🟥 Đỏ: chỉ có trong file 1")
    print("🟦 Xanh dương: có trong file 1 & xuất hiện >1 lần trong file 2")
    print("🟨 Vàng: chỉ có trong file 2")


# --- Ví dụ chạy ---
if __name__ == "__main__":
    compare_excel_fast("1.xlsx", "2.xlsx", column_name="Name")
