import polars as pl
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment

def compare_excel_fast(file1, file2, column_name, output_file="compare_result_polars.xlsx"):
    # --- Đọc file bằng Polars (đa luồng, rất nhanh) ---
    df1 = pl.read_excel(file1)
    df2 = pl.read_excel(file2)

    # --- Kiểm tra cột tồn tại ---
    if column_name not in df1.columns or column_name not in df2.columns:
        raise ValueError(f"Cột '{column_name}' không tồn tại trong một trong hai file.")

    # --- Ép kiểu về string ---
    df1 = df1.with_columns(pl.col(column_name).cast(pl.Utf8))
    df2 = df2.with_columns(pl.col(column_name).cast(pl.Utf8))

    # --- Tạo tập giá trị ---
    set1 = set(df1[column_name].to_list())
    set2 = set(df2[column_name].to_list())
    common_keys = set1 & set2
    only_in_file1 = set1 - set2
    only_in_file2 = set2 - set1

    # --- Đếm xuất hiện trong file 2 ---
    duplicated_in_file2 = set(
        df2[column_name].value_counts()
        .filter(pl.col("count") > 1)[column_name]
        .to_list()
    ) & set1

    # --- Gán trạng thái ---
    def assign_status(df, file_index):
        def label(v):
            if file_index == 1:
                if v in only_in_file1:
                    return "red"
                elif v in duplicated_in_file2:
                    return "blue"
                elif v in common_keys:
                    return "green"
            else:
                if v in only_in_file2:
                    return "yellow"
                elif v in duplicated_in_file2:
                    return "blue"
                elif v in common_keys:
                    return "green"
            return ""
        return df.with_columns(
            pl.col(column_name).map_elements(label, return_dtype=pl.Utf8).alias("__status__")
        )

    df1 = assign_status(df1, 1)
    df2 = assign_status(df2, 2)

    # --- Chuyển sang Pandas để ghi Excel ---
    pd_df1 = df1.to_pandas()
    pd_df2 = df2.to_pandas()

    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        pd_df1.to_excel(writer, sheet_name="File1", index=False)
        pd_df2.to_excel(writer, sheet_name="File2", index=False)

    # --- Mở Excel và tô màu ---
    wb = load_workbook(output_file)
    fills = {
        "green":  PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
        "red":    PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
        "blue":   PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid"),
        "yellow": PatternFill(start_color="FFFACD", end_color="FFFACD", fill_type="solid"),
    }

    for sheet_name in ["File1", "File2"]:
        ws = wb[sheet_name]
        status_col = ws.max_column
        header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        header_font = Font(bold=True)
        header_align = Alignment(horizontal="center", vertical="center")

        # Tô header
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align

        # Tô màu theo trạng thái
        for r in range(2, ws.max_row + 1):
            color_key = ws.cell(row=r, column=status_col).value
            if color_key in fills:
                for c in range(1, ws.max_column + 1):
                    ws.cell(row=r, column=c).fill = fills[color_key]

        # Xóa cột "__status__"
        ws.delete_cols(status_col)

        # Căn chỉnh cột đẹp hơn
        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value else 0) for c in col)
            ws.column_dimensions[col[0].column_letter].width = max_len + 2

    wb.save(output_file)

    print(f"✅ Đã tạo file '{output_file}' thành công.")
    print("🟩 Xanh lá: có trong cả hai file")
    print("🟥 Đỏ: chỉ có trong file 1")
    print("🟦 Xanh dương: có trong file 1 và xuất hiện >1 lần trong file 2")
    print("🟨 Vàng: chỉ có trong file 2")

# --- Ví dụ chạy ---
compare_excel_fast("payments-Copy.xlsx", "transactions - Copy.xlsx", column_name="Payments_id")
